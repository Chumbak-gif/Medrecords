"""Medicines router — thin controller delegating to use cases.

Preserves all existing endpoint paths, methods, query params, and response shapes.
"""

import io
from typing import Optional

from fastapi import APIRouter, Depends, File, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, field_validator
from fastapi import HTTPException
from datetime import datetime

from src.api.v1.dependencies.auth import get_current_user, role_required
from src.api.v1.dependencies.container import (
    get_create_medicine_use_case,
    get_delete_medicine_use_case,
    get_get_medicine_use_case,
    get_import_medicines_use_case,
    get_list_medicines_use_case,
    get_update_medicine_use_case,
)
from src.api.v1.schemas.common import PaginatedResponse
from src.application.use_cases.medicines.create_medicine import CreateMedicineCommand, CreateMedicineUseCase
from src.application.use_cases.medicines.delete_medicine import DeleteMedicineCommand, DeleteMedicineUseCase
from src.application.use_cases.medicines.get_medicine import GetMedicineQuery, GetMedicineUseCase
from src.application.use_cases.medicines.import_medicines import ImportMedicineEntry, ImportMedicinesCommand, ImportMedicinesUseCase
from src.application.use_cases.medicines.list_medicines import ListMedicinesQuery, ListMedicinesUseCase
from src.application.use_cases.medicines.update_medicine import UpdateMedicineCommand, UpdateMedicineUseCase
from src.domain.entities.user import UserEntity
from src.infrastructure.persistence.models.user_model import UserModel

router = APIRouter(tags=["Medicines"])


# --- Schemas ---

class MedicineCreate(BaseModel):
    name: str
    brand_name: Optional[str] = None
    generic_name: Optional[str] = None
    strength: Optional[str] = None
    form: Optional[str] = None
    manufacturer: Optional[str] = None
    category: Optional[str] = None
    unit: Optional[str] = None

    @field_validator("name")
    @classmethod
    def name_must_not_be_blank(cls, v: str) -> str:
        v = v.strip()
        if not v:
            raise ValueError("name must not be blank")
        return v


class MedicineUpdate(BaseModel):
    name: Optional[str] = None
    brand_name: Optional[str] = None
    generic_name: Optional[str] = None
    strength: Optional[str] = None
    form: Optional[str] = None
    manufacturer: Optional[str] = None
    category: Optional[str] = None
    unit: Optional[str] = None
    is_active: Optional[bool] = None

    @field_validator("name")
    @classmethod
    def name_must_not_be_blank(cls, v: Optional[str]) -> Optional[str]:
        if v is not None:
            v = v.strip()
            if not v:
                raise ValueError("name must not be blank")
        return v


class MedicineResponse(BaseModel):
    id: int
    name: str
    brand_name: Optional[str] = None
    generic_name: Optional[str] = None
    strength: Optional[str] = None
    form: Optional[str] = None
    manufacturer: Optional[str] = None
    category: Optional[str] = None
    unit: Optional[str] = None
    is_active: bool
    created_at: datetime
    updated_at: datetime

    model_config = {"from_attributes": True}


class ImportErrorItem(BaseModel):
    row: int
    field: str
    message: str


class ImportResultResponse(BaseModel):
    inserted: int
    updated: int
    errors: list[ImportErrorItem]


# --- Helpers ---

def _to_actor(user_model: UserModel) -> UserEntity:
    return UserEntity(
        id=user_model.id,
        username=user_model.username,
        email=user_model.email,
        full_name=user_model.full_name,
        hashed_password=user_model.hashed_password,
        role=user_model.role,
        specialty=user_model.specialty,
        is_active=user_model.is_active,
        created_at=user_model.created_at,
        updated_at=user_model.updated_at,
    )


# --- Fixed-path routes (before /{id}) ---

@router.post("/import", response_model=ImportResultResponse, summary="Bulk import medicines from an .xlsx file (admin only)")
async def import_medicines(
    file: UploadFile = File(...),
    current_user: UserModel = Depends(role_required(["admin", "sys_admin"])),
    use_case: ImportMedicinesUseCase = Depends(get_import_medicines_use_case),
) -> ImportResultResponse:
    contents = await file.read()
    try:
        wb = load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read the uploaded file. Please upload a valid .xlsx file.")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded file is empty.")

    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    try:
        name_col = header_row.index("name")
    except ValueError:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing required column 'name' in the header row.")

    category_col: Optional[int] = header_row.index("category") if "category" in header_row else None
    unit_col: Optional[int] = header_row.index("unit") if "unit" in header_row else None
    brand_name_col: Optional[int] = header_row.index("brand_name") if "brand_name" in header_row else None
    generic_name_col: Optional[int] = header_row.index("generic_name") if "generic_name" in header_row else None
    strength_col: Optional[int] = header_row.index("strength") if "strength" in header_row else None
    form_col: Optional[int] = header_row.index("form") if "form" in header_row else None
    manufacturer_col: Optional[int] = header_row.index("manufacturer") if "manufacturer" in header_row else None

    data_rows = rows[1:]
    errors: list[ImportErrorItem] = []
    parsed: list[ImportMedicineEntry] = []

    for row_idx, row in enumerate(data_rows, start=2):
        raw_name = row[name_col] if len(row) > name_col else None
        name = str(raw_name).strip() if raw_name is not None else ""

        if not name:
            errors.append(ImportErrorItem(row=row_idx, field="name", message="name is required and must not be blank"))
            continue

        def _get_col(col: Optional[int]) -> Optional[str]:
            if col is None or len(row) <= col:
                return None
            val = row[col]
            if val is None:
                return None
            s = str(val).strip()
            return s if s else None

        parsed.append(ImportMedicineEntry(
            name=name, brand_name=_get_col(brand_name_col), generic_name=_get_col(generic_name_col),
            strength=_get_col(strength_col), form=_get_col(form_col), manufacturer=_get_col(manufacturer_col),
            category=_get_col(category_col), unit=_get_col(unit_col),
        ))

    if errors:
        return ImportResultResponse(inserted=0, updated=0, errors=errors)

    command = ImportMedicinesCommand(entries=parsed, actor=_to_actor(current_user))
    result = await use_case.execute(command)
    return ImportResultResponse(inserted=result.inserted, updated=result.updated, errors=[])


@router.get("/template", summary="Download blank .xlsx import template (admin only)")
async def download_import_template(_current_user: UserModel = Depends(role_required(["admin", "sys_admin"]))) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Medicines"
    ws.append(["name", "brand_name", "generic_name", "strength", "form", "manufacturer", "category", "unit"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=medicines_template.xlsx"})


# --- CRUD endpoints ---

@router.get("/", response_model=PaginatedResponse[MedicineResponse], summary="List medicines (paginated + search)")
async def list_medicines(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: str = Query(""),
    include_inactive: bool = Query(False),
    _current_user: UserModel = Depends(get_current_user),
    use_case: ListMedicinesUseCase = Depends(get_list_medicines_use_case),
) -> PaginatedResponse[MedicineResponse]:
    query = ListMedicinesQuery(
        page=page,
        page_size=page_size,
        search=search,
        include_inactive=include_inactive,
    )
    result = await use_case.execute(query)
    return PaginatedResponse(
        items=[MedicineResponse.model_validate(m) for m in result.items],
        total=result.total,
        page=result.page,
        page_size=result.page_size,
        total_pages=result.total_pages,
    )


@router.post("/", response_model=MedicineResponse, status_code=status.HTTP_201_CREATED, summary="Create a new medicine (admin only)")
async def create_medicine(
    payload: MedicineCreate,
    current_user: UserModel = Depends(role_required(["admin", "sys_admin"])),
    use_case: CreateMedicineUseCase = Depends(get_create_medicine_use_case),
) -> MedicineResponse:
    command = CreateMedicineCommand(
        name=payload.name,
        brand_name=payload.brand_name,
        generic_name=payload.generic_name,
        strength=payload.strength,
        form=payload.form,
        manufacturer=payload.manufacturer,
        category=payload.category,
        unit=payload.unit,
        actor=_to_actor(current_user),
    )
    medicine = await use_case.execute(command)
    return MedicineResponse.model_validate(medicine)


@router.get("/{id}", response_model=MedicineResponse, summary="Get a single medicine by ID")
async def get_medicine(
    id: int,
    _current_user: UserModel = Depends(get_current_user),
    use_case: GetMedicineUseCase = Depends(get_get_medicine_use_case),
) -> MedicineResponse:
    query = GetMedicineQuery(medicine_id=id)
    medicine = await use_case.execute(query)
    return MedicineResponse.model_validate(medicine)


@router.patch("/{id}", response_model=MedicineResponse, summary="Update a medicine (admin only)")
async def update_medicine(
    id: int,
    payload: MedicineUpdate,
    current_user: UserModel = Depends(role_required(["admin", "sys_admin"])),
    use_case: UpdateMedicineUseCase = Depends(get_update_medicine_use_case),
) -> MedicineResponse:
    command = UpdateMedicineCommand(
        medicine_id=id,
        name=payload.name,
        brand_name=payload.brand_name,
        generic_name=payload.generic_name,
        strength=payload.strength,
        form=payload.form,
        manufacturer=payload.manufacturer,
        category=payload.category,
        unit=payload.unit,
        is_active=payload.is_active,
        actor=_to_actor(current_user),
    )
    medicine = await use_case.execute(command)
    return MedicineResponse.model_validate(medicine)


@router.delete("/{id}", status_code=status.HTTP_204_NO_CONTENT, summary="Soft-delete a medicine (admin only)")
async def delete_medicine(
    id: int,
    current_user: UserModel = Depends(role_required(["admin", "sys_admin"])),
    use_case: DeleteMedicineUseCase = Depends(get_delete_medicine_use_case),
) -> None:
    command = DeleteMedicineCommand(
        medicine_id=id,
        actor=_to_actor(current_user),
    )
    await use_case.execute(command)
