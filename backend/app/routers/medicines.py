"""
Medicine router
Prefix:  /api/v1/medicines  (mounted in main.py)

Route registration order is important — /import and /template MUST come
before /{id} so that FastAPI does not attempt to match those literal path
segments as integer IDs.
"""
import io
import math
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.dependencies.auth import get_current_user, role_required
from app.models.audit_log import AuditLog
from app.models.medicine import Medicine
from app.models.user import User
from app.schemas.disease import PaginatedResponse
from app.schemas.medicine import (
    ImportError,
    ImportResult,
    MedicineCreate,
    MedicineResponse,
    MedicineUpdate,
)

router = APIRouter(tags=["Medicines"])

# ---------------------------------------------------------------------------
# Internal helper — write an audit log entry
# ---------------------------------------------------------------------------

async def _audit(
    db: AsyncSession,
    *,
    event_type: str,
    actor: User,
    entity_type: str,
    entity_id: int,
    description: str,
) -> None:
    db.add(
        AuditLog(
            event_type=event_type,
            actor_id=actor.id,
            actor_username=actor.username,
            actor_role=actor.role,
            entity_type=entity_type,
            entity_id=entity_id,
            description=description,
        )
    )


# ===========================================================================
# Fixed-path routes — registered BEFORE /{id} to avoid routing clashes
# ===========================================================================

@router.post(
    "/import",
    response_model=ImportResult,
    summary="Bulk import medicines from an .xlsx file (admin only)",
)
async def import_medicines(
    file: UploadFile = File(..., description="Excel (.xlsx) file with columns: name, category, unit"),
    current_user: User = Depends(role_required(["admin", "sys_admin"])),
    db: AsyncSession = Depends(get_db),
) -> ImportResult:
    # ── 1. Read uploaded bytes ──────────────────────────────────────────────
    contents = await file.read()
    try:
        wb = load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not read the uploaded file. Please upload a valid .xlsx file.",
        )

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty.",
        )

    # ── 2. Resolve header row ───────────────────────────────────────────────
    header_row = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    try:
        name_col = header_row.index("name")
    except ValueError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Missing required column 'name' in the header row.",
        )

    category_col: Optional[int] = header_row.index("category") if "category" in header_row else None
    unit_col: Optional[int] = header_row.index("unit") if "unit" in header_row else None
    brand_name_col: Optional[int] = header_row.index("brand_name") if "brand_name" in header_row else None
    generic_name_col: Optional[int] = header_row.index("generic_name") if "generic_name" in header_row else None
    strength_col: Optional[int] = header_row.index("strength") if "strength" in header_row else None
    form_col: Optional[int] = header_row.index("form") if "form" in header_row else None
    manufacturer_col: Optional[int] = header_row.index("manufacturer") if "manufacturer" in header_row else None

    data_rows = rows[1:]

    # ── 3. Validate ALL rows first — collect errors without aborting early ──
    errors: list[ImportError] = []
    parsed: list[dict] = []  # holds cleaned data for rows that pass validation

    for row_idx, row in enumerate(data_rows, start=2):  # row 1 is header; data starts at 2
        raw_name = row[name_col] if len(row) > name_col else None
        name = str(raw_name).strip() if raw_name is not None else ""

        if not name:
            errors.append(ImportError(row=row_idx, field="name", message="name is required and must not be blank"))
            continue

        def _get_col(col: Optional[int]) -> Optional[str]:
            if col is None or len(row) <= col:
                return None
            val = row[col]
            if val is None:
                return None
            s = str(val).strip()
            return s if s else None

        parsed.append({
            "name": name,
            "brand_name": _get_col(brand_name_col),
            "generic_name": _get_col(generic_name_col),
            "strength": _get_col(strength_col),
            "form": _get_col(form_col),
            "manufacturer": _get_col(manufacturer_col),
            "category": _get_col(category_col),
            "unit": _get_col(unit_col),
        })

    # ── 4. Reject entire batch if ANY row has errors ─────────────────────────
    if errors:
        return ImportResult(inserted=0, updated=0, errors=errors)

    # ── 5. Upsert — update existing by name, insert new ──────────────────────
    inserted = 0
    updated = 0

    for entry in parsed:
        result = await db.execute(
            select(Medicine).where(func.lower(Medicine.name) == entry["name"].lower())
        )
        existing: Optional[Medicine] = result.scalar_one_or_none()

        if existing is not None:
            existing.brand_name = entry["brand_name"]
            existing.generic_name = entry["generic_name"]
            existing.strength = entry["strength"]
            existing.form = entry["form"]
            existing.manufacturer = entry["manufacturer"]
            existing.category = entry["category"]
            existing.unit = entry["unit"]
            updated += 1
        else:
            db.add(Medicine(
                name=entry["name"],
                brand_name=entry["brand_name"],
                generic_name=entry["generic_name"],
                strength=entry["strength"],
                form=entry["form"],
                manufacturer=entry["manufacturer"],
                category=entry["category"],
                unit=entry["unit"],
            ))
            inserted += 1

    await db.flush()

    await _audit(
        db,
        event_type="medicine_bulk_import",
        actor=current_user,
        entity_type="medicine",
        entity_id=0,
        description=f"Bulk import: {inserted} inserted, {updated} updated",
    )

    return ImportResult(inserted=inserted, updated=updated, errors=[])


@router.get(
    "/template",
    summary="Download blank .xlsx import template (admin only)",
)
async def download_import_template(
    _current_user: User = Depends(role_required(["admin", "sys_admin"])),
) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Medicines"
    ws.append(["name", "brand_name", "generic_name", "strength", "form", "manufacturer", "category", "unit"])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=medicines_template.xlsx"},
    )


# ===========================================================================
# CRUD endpoints
# ===========================================================================

@router.get(
    "/",
    response_model=PaginatedResponse[MedicineResponse],
    summary="List medicines (paginated + search)",
)
async def list_medicines(
    page: int = Query(1, ge=1, description="Page number (1-based)"),
    page_size: int = Query(20, ge=1, le=100, description="Items per page"),
    search: str = Query("", description="Search term matched against name"),
    include_inactive: bool = Query(False, description="Include soft-deleted medicines"),
    _current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> PaginatedResponse[MedicineResponse]:
    q = select(Medicine)
    if not include_inactive:
        q = q.where(Medicine.is_active.is_(True))
    if search:
        q = q.where(Medicine.name.ilike(f"%{search}%"))

    count_q = select(func.count()).select_from(q.subquery())
    total: int = (await db.execute(count_q)).scalar_one()

    offset = (page - 1) * page_size
    result = await db.execute(q.order_by(Medicine.name).offset(offset).limit(page_size))
    medicines = result.scalars().all()

    return PaginatedResponse(
        items=medicines,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=max(1, math.ceil(total / page_size)),
    )


@router.post(
    "/",
    response_model=MedicineResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Create a new medicine (admin only)",
)
async def create_medicine(
    payload: MedicineCreate,
    current_user: User = Depends(role_required(["admin", "sys_admin"])),
    db: AsyncSession = Depends(get_db),
) -> MedicineResponse:
    existing = await db.execute(
        select(Medicine).where(func.lower(Medicine.name) == payload.name.lower())
    )
    if existing.scalar_one_or_none():
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"A medicine named '{payload.name}' already exists",
        )

    medicine = Medicine(
        name=payload.name,
        brand_name=payload.brand_name,
        generic_name=payload.generic_name,
        strength=payload.strength,
        form=payload.form,
        manufacturer=payload.manufacturer,
        category=payload.category,
        unit=payload.unit,
    )
    db.add(medicine)
    await db.flush()

    await _audit(
        db,
        event_type="medicine_created",
        actor=current_user,
        entity_type="medicine",
        entity_id=medicine.id,
        description=f"Medicine '{medicine.name}' created",
    )

    return medicine


@router.get(
    "/{id}",
    response_model=MedicineResponse,
    summary="Get a single medicine by ID",
)
async def get_medicine(
    id: int,
    _current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> MedicineResponse:
    result = await db.execute(select(Medicine).where(Medicine.id == id))
    medicine: Optional[Medicine] = result.scalar_one_or_none()
    if medicine is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Medicine not found")
    return medicine


@router.patch(
    "/{id}",
    response_model=MedicineResponse,
    summary="Update a medicine (admin only)",
)
async def update_medicine(
    id: int,
    payload: MedicineUpdate,
    current_user: User = Depends(role_required(["admin", "sys_admin"])),
    db: AsyncSession = Depends(get_db),
) -> MedicineResponse:
    result = await db.execute(select(Medicine).where(Medicine.id == id))
    medicine: Optional[Medicine] = result.scalar_one_or_none()
    if medicine is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Medicine not found")

    if payload.name is not None and payload.name.lower() != medicine.name.lower():
        existing = await db.execute(
            select(Medicine).where(
                func.lower(Medicine.name) == payload.name.lower(),
                Medicine.id != id,
            )
        )
        if existing.scalar_one_or_none():
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"A medicine named '{payload.name}' already exists",
            )
        medicine.name = payload.name

    if payload.category is not None:
        medicine.category = payload.category
    if payload.unit is not None:
        medicine.unit = payload.unit
    if payload.brand_name is not None:
        medicine.brand_name = payload.brand_name
    if payload.generic_name is not None:
        medicine.generic_name = payload.generic_name
    if payload.strength is not None:
        medicine.strength = payload.strength
    if payload.form is not None:
        medicine.form = payload.form
    if payload.manufacturer is not None:
        medicine.manufacturer = payload.manufacturer
    if payload.is_active is not None:
        medicine.is_active = payload.is_active

    await _audit(
        db,
        event_type="medicine_updated",
        actor=current_user,
        entity_type="medicine",
        entity_id=medicine.id,
        description=f"Medicine '{medicine.name}' updated",
    )

    return medicine


@router.delete(
    "/{id}",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Soft-delete a medicine (admin only)",
)
async def delete_medicine(
    id: int,
    current_user: User = Depends(role_required(["admin", "sys_admin"])),
    db: AsyncSession = Depends(get_db),
) -> None:
    result = await db.execute(select(Medicine).where(Medicine.id == id))
    medicine: Optional[Medicine] = result.scalar_one_or_none()
    if medicine is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Medicine not found")

    medicine.is_active = False

    await _audit(
        db,
        event_type="medicine_deleted",
        actor=current_user,
        entity_type="medicine",
        entity_id=medicine.id,
        description=f"Medicine '{medicine.name}' soft-deleted",
    )
